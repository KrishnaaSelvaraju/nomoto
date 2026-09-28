"""
MSMD Assignment 2026/2027 - Project 9, Task 2
Head-on encounter: COLREG-COLREG and Blind-Blind, using corrected T.
"""

import json
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.animation as animation
from matplotlib.patches import Polygon

import ship_data as sd
from engine import find_minimum_safe_distance

SEED = 2026  # for Blind-Blind's representative rudder draw


def matrices_to_excel(rows1, rows2, cols, D_min, scenario_name, out_path):
    df1 = pd.DataFrame(rows1, columns=cols)
    df2 = pd.DataFrame(rows2, columns=cols)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df1.to_excel(writer, sheet_name=f"{sd.VESSEL1['name']} (Vessel 1)", index=False)
        df2.to_excel(writer, sheet_name=f"{sd.VESSEL2['name']} (Vessel 2)", index=False)
        summary = pd.DataFrame({
            "Parameter": ["Scenario", "Minimum safe distance (m)",
                          "Minimum safe distance / larger Lpp"],
            "Value": [scenario_name, round(D_min, 1),
                      round(D_min / max(sd.VESSEL1["Lpp"], sd.VESSEL2["Lpp"]), 3)],
        })
        summary.to_excel(writer, sheet_name="Summary", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(out_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", bold=(cell.row == 1))
        for col_cells in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 22)
    wb.save(out_path)


def plot_encounter(rows1, rows2, D_min, scenario_name, out_path):
    fig, ax = plt.subplots(figsize=(7, 8))
    ax.plot(rows1[:, 1], rows1[:, 2], color="tab:blue", label=f"{sd.VESSEL1['name']} path")
    ax.plot(rows2[:, 1], rows2[:, 2], color="tab:orange", label=f"{sd.VESSEL2['name']} path")
    for rows, color in [(rows1, "tab:blue"), (rows2, "tab:orange")]:
        for frac in (0.0, 0.5, 1.0):
            i = int(frac * (len(rows) - 1))
            corners = rows[i, 3:11].reshape(4, 2)
            poly = plt.Polygon(corners[[0, 1, 3, 2]], closed=True, facecolor=color,
                                edgecolor="black", alpha=0.4)
            ax.add_patch(poly)
    ax.set_xlabel("x (m)"); ax.set_ylabel("y (m)")
    ax.set_title(f"Head-on, {scenario_name}: minimum safe distance = {D_min:.0f} m")
    ax.axis("equal"); ax.grid(alpha=0.3); ax.legend()
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def ship_polygon_local(length, beam):
    l2, b2 = length / 2.0, beam / 2.0
    return np.array([[0.0, l2], [b2, l2 * 0.6], [b2, -l2], [-b2, -l2], [-b2, l2 * 0.6]])


def animate_encounter(rows1, rows2, scenario_name, out_path, dt, fps=25, speed_up=4):
    step = max(1, int(round(speed_up / dt / fps)))
    n = min(len(rows1), len(rows2))
    idx = np.arange(0, n, step)

    fig, ax = plt.subplots(figsize=(7, 8))
    ax.plot(rows1[:, 1], rows1[:, 2], color="lightblue", linewidth=1.2, zorder=1)
    ax.plot(rows2[:, 1], rows2[:, 2], color="navajowhite", linewidth=1.2, zorder=1)
    all_x = np.concatenate([rows1[:, 1], rows2[:, 1]])
    all_y = np.concatenate([rows1[:, 2], rows2[:, 2]])
    margin = 0.1 * max(all_x.max() - all_x.min(), all_y.max() - all_y.min(), 1.0)
    ax.set_xlim(all_x.min() - margin, all_x.max() + margin)
    ax.set_ylim(all_y.min() - margin, all_y.max() + margin)
    ax.set_aspect("equal"); ax.set_xlabel("x (m)"); ax.set_ylabel("y (m)"); ax.grid(alpha=0.3)
    title = ax.set_title("")

    poly1_local = ship_polygon_local(sd.VESSEL1["Lpp"], sd.VESSEL1["B"])
    poly2_local = ship_polygon_local(sd.VESSEL2["Lpp"], sd.VESSEL2["B"])
    patch1 = Polygon(poly1_local, closed=True, facecolor="steelblue", edgecolor="black",
                      zorder=3, label=sd.VESSEL1["name"])
    patch2 = Polygon(poly2_local, closed=True, facecolor="darkorange", edgecolor="black",
                      zorder=3, label=sd.VESSEL2["name"])
    ax.add_patch(patch1); ax.add_patch(patch2)
    ax.legend(handles=[patch1, patch2], loc="upper right")

    def rotate_translate(local_poly, psi_deg, x, y):
        psi = np.radians(psi_deg)
        c, s = np.cos(psi), np.sin(psi)
        rot = np.array([[c, s], [-s, c]])
        return local_poly @ rot.T + np.array([x, y])

    def init():
        patch1.set_xy(poly1_local); patch2.set_xy(poly2_local)
        return patch1, patch2, title

    def update(frame_i):
        i = idx[frame_i]
        patch1.set_xy(rotate_translate(poly1_local, rows1[i, 12], rows1[i, 1], rows1[i, 2]))
        patch2.set_xy(rotate_translate(poly2_local, rows2[i, 12], rows2[i, 1], rows2[i, 2]))
        title.set_text(f"Head-on, {scenario_name} - t = {rows1[i, 0]:.0f} s")
        return patch1, patch2, title

    ani = animation.FuncAnimation(fig, update, frames=len(idx), init_func=init, blit=False,
                                   interval=1000 / fps)
    writer = animation.FFMpegWriter(fps=fps, bitrate=2400)
    ani.save(out_path, writer=writer)
    plt.close(fig)


if __name__ == "__main__":
    results = {}

    print("=== Head-on, COLREG-COLREG ===")
    print("Rule 14: both vessels alter course to starboard, delta = 20 deg (Tri most-likely).")
    D1, r1a, r1b, cols = find_minimum_safe_distance(delta1_target_deg=+20.0, delta2_target_deg=+20.0)
    results["COLREG-COLREG"] = {"D_min": D1, "rows1": r1a, "rows2": r1b}

    print("\n=== Head-on, Blind-Blind ===")
    rng = np.random.default_rng(SEED)
    delta1_blind = rng.uniform(-35, 35)
    delta2_blind = rng.uniform(-35, 35)
    print(f"Sampled blind rudder angles (seed={SEED}): "
          f"delta1={delta1_blind:.1f} deg, delta2={delta2_blind:.1f} deg")
    D2, r2a, r2b, _ = find_minimum_safe_distance(delta1_target_deg=delta1_blind,
                                                  delta2_target_deg=delta2_blind)
    results["Blind-Blind"] = {"D_min": D2, "rows1": r2a, "rows2": r2b,
                               "delta1_deg": delta1_blind, "delta2_deg": delta2_blind}

    summary_json = {}
    for name, res in results.items():
        tag = name.lower().replace("-", "_")
        D, rows1, rows2 = res["D_min"], res["rows1"], res["rows2"]
        print(f"\n{name}: minimum safe distance = {D:.1f} m")

        matrices_to_excel(rows1, rows2, cols, D, name, f"outputs/task2_headon_{tag}_data.xlsx")
        plot_encounter(rows1, rows2, D, name, f"figures/task2_headon_{tag}.png")
        animate_encounter(rows1, rows2, name, f"figures/task2_headon_{tag}.mp4", dt=0.5)
        print(f"  Saved data table, plot, and MP4 for {name}")

        summary_json[name] = {
            "D_min_m": float(D),
            "D_min_over_larger_Lpp": float(D / max(sd.VESSEL1["Lpp"], sd.VESSEL2["Lpp"])),
        }
        if "delta1_deg" in res:
            summary_json[name]["delta1_deg"] = float(res["delta1_deg"])
            summary_json[name]["delta2_deg"] = float(res["delta2_deg"])

    with open("outputs/task2_results.json", "w") as f:
        json.dump(summary_json, f, indent=2)
    print("\nSaved outputs/task2_results.json")
